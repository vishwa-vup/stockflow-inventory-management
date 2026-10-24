"""Build the StockFlow project report as a Word document.

Structure and styling follow the supplied Medicine Schedule and Reminder Manager
report: A4, Times New Roman, front matter with roman numerals, chapters from 1,
captioned tables and figures.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUTPUT = HERE / "StockFlow_Project_Report.docx"

BLACK = RGBColor(0, 0, 0)
NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x47, 0x55, 0x69)
WIDTH = 6.2  # usable text width in inches

TEAM = [
    ("SAI VARSHINI J", "25AM104"),
    ("SAPARNA S", "25AM105"),
    ("SASMITHA K M", "25AM106"),
    ("SENAAN Y A", "25AM107"),
]
PROJECT = "StockFlow - Inventory and Stock Management System"
PROJECT_SHORT = "StockFlow"


class Report:
    def __init__(self):
        self.doc = Document()
        self._page_setup()
        self._styles()

    # ------------------------------------------------------------------ setup
    def _page_setup(self):
        section = self.doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.1)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.45)
        section.footer_distance = Inches(0.45)

        props = self.doc.core_properties
        props.title = "StockFlow - Inventory and Stock Management System"
        props.subject = "Project Report"
        props.author = ", ".join(name for name, _ in TEAM)
        props.comments = "Bachelor of Engineering, CSE (Artificial Intelligence and Machine Learning)"

    def _set_font(self, style, size, bold=False, italic=False, color=BLACK):
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = color
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), "Times New Roman")

    def _styles(self):
        normal = self.doc.styles["Normal"]
        self._set_font(normal, 12)
        pf = normal.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.widow_control = True

        h1 = self.doc.styles["Heading 1"]
        self._set_font(h1, 16, bold=True)
        h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h1.paragraph_format.space_before = Pt(0)
        h1.paragraph_format.space_after = Pt(14)
        h1.paragraph_format.line_spacing = 1.15
        h1.paragraph_format.keep_with_next = True

        h3 = self.doc.styles["Heading 3"]
        self._set_font(h3, 13, bold=True)
        h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h3.paragraph_format.space_before = Pt(14)
        h3.paragraph_format.space_after = Pt(4)
        h3.paragraph_format.line_spacing = 1.15
        h3.paragraph_format.keep_with_next = True

        h4 = self.doc.styles["Heading 4"]
        self._set_font(h4, 12, bold=True)
        h4.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h4.paragraph_format.space_before = Pt(10)
        h4.paragraph_format.space_after = Pt(3)
        h4.paragraph_format.line_spacing = 1.15
        h4.paragraph_format.keep_with_next = True

        for name in ("List Bullet", "List Number"):
            st = self.doc.styles[name]
            self._set_font(st, 12)
            st.paragraph_format.line_spacing = 1.35
            st.paragraph_format.space_after = Pt(3)

        fig = self.doc.styles.add_style("FigCaption", WD_STYLE_TYPE.PARAGRAPH)
        self._set_font(fig, 10.5, italic=True, color=GREY)
        fig.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fig.paragraph_format.space_before = Pt(2)
        fig.paragraph_format.space_after = Pt(14)
        fig.paragraph_format.line_spacing = 1.0

        cap = self.doc.styles.add_style("TableCaption", WD_STYLE_TYPE.PARAGRAPH)
        self._set_font(cap, 10.5, bold=True, color=NAVY)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(8)
        cap.paragraph_format.space_after = Pt(4)
        cap.paragraph_format.line_spacing = 1.0
        cap.paragraph_format.keep_with_next = True

    # ---------------------------------------------------------------- helpers
    def p(self, text="", lead=None, style=None, align=None, italic=False,
          size=None, bold=False, space_after=None, space_before=None):
        par = self.doc.add_paragraph(style=style)
        if lead:
            run = par.add_run(lead)
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        run = par.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = "Times New Roman"
        if size:
            run.font.size = Pt(size)
        if align is not None:
            par.alignment = align
        if space_after is not None:
            par.paragraph_format.space_after = Pt(space_after)
        if space_before is not None:
            par.paragraph_format.space_before = Pt(space_before)
        return par

    def runs(self, parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12,
             space_after=6, line_spacing=1.5):
        par = self.doc.add_paragraph()
        par.alignment = align
        par.paragraph_format.space_after = Pt(space_after)
        par.paragraph_format.line_spacing = line_spacing
        for text, bold in parts:
            run = par.add_run(text)
            run.bold = bold
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
        return par

    def bullets(self, items, style="List Bullet"):
        # Keep the lead-in sentence on the same page as the first bullets.
        if self.doc.paragraphs:
            prev = self.doc.paragraphs[-1]
            if prev.style.name not in ("List Bullet", "List Number"):
                prev.paragraph_format.keep_with_next = True
        for item in items:
            par = self.doc.add_paragraph(style=style)
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.left_indent = Inches(0.35)
            par.paragraph_format.first_line_indent = Inches(-0.2)
            run = par.add_run(item)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

    def h1(self, text, page_break=True):
        par = self.doc.add_paragraph(text, style="Heading 1")
        par.paragraph_format.page_break_before = page_break
        return par

    def h3(self, text):
        return self.doc.add_paragraph(text, style="Heading 3")

    def h4(self, text):
        return self.doc.add_paragraph(text, style="Heading 4")

    def figure(self, filename, width, caption):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(8)
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
        cap = self.doc.add_paragraph(caption, style="FigCaption")
        return cap

    def _shade(self, cell, fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    def _margins(self, cell, top=50, bottom=50, left=90, right=90):
        tc_pr = cell._tc.get_or_add_tcPr()
        mar = OxmlElement("w:tcMar")
        for tag, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
            node = OxmlElement(f"w:{tag}")
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")
            mar.append(node)
        tc_pr.append(mar)

    def _write_cell(self, cell, text, bold=False, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
        cell.text = ""
        par = cell.paragraphs[0]
        par.alignment = align
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        run = par.add_run(str(text))
        run.bold = bold
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)

    def table(self, caption, headers, rows, widths, font=9.5,
              header_fill="D9E2F3", center_cols=(), first_col_bold=False):
        if caption:
            self.doc.add_paragraph(caption, style="TableCaption")
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        tbl_pr = table._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tbl_pr.append(layout)

        for index, width in enumerate(widths):
            table.columns[index].width = Inches(width)

        header = table.rows[0]
        tr_pr = header._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)
        for index, text in enumerate(headers):
            cell = header.cells[index]
            cell.width = Inches(widths[index])
            self._write_cell(cell, text, bold=True, size=font)
            self._shade(cell, header_fill)
            self._margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        for r_index, row in enumerate(rows):
            cells = table.add_row().cells
            tr = table.rows[-1]._tr
            tr_pr = tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            tr_pr.append(cant)
            for c_index, text in enumerate(row):
                cell = cells[c_index]
                cell.width = Inches(widths[c_index])
                align = WD_ALIGN_PARAGRAPH.CENTER if c_index in center_cols else WD_ALIGN_PARAGRAPH.LEFT
                self._write_cell(cell, text, bold=(first_col_bold and c_index == 0),
                                 size=font, align=align)
                self._margins(cell)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if r_index % 2 == 1:
                    self._shade(cell, "F5F8FC")

        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        spacer.paragraph_format.line_spacing = 1.0
        return table

    def code(self, lines, caption=None):
        if caption:
            self.doc.add_paragraph(caption, style="TableCaption")
        table = self.doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        cell.width = Inches(WIDTH)
        self._shade(cell, "F4F5F7")
        self._margins(cell, top=100, bottom=100, left=140, right=140)
        cell.text = ""
        first = True
        for line in lines:
            par = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            run = par.add_run(line)
            run.font.name = "Courier New"
            run.font.size = Pt(8.5)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        spacer.paragraph_format.line_spacing = 1.0

    def _page_field(self, paragraph):
        run = paragraph.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(instr)
        run._r.append(end)

    def _numbering(self, section, fmt, start):
        sect_pr = section._sectPr
        for old in sect_pr.findall(qn("w:pgNumType")):
            sect_pr.remove(old)
        node = OxmlElement("w:pgNumType")
        node.set(qn("w:fmt"), fmt)
        node.set(qn("w:start"), str(start))
        cols = sect_pr.find(qn("w:cols"))
        if cols is not None:
            cols.addprevious(node)
        else:
            sect_pr.append(node)

    def _footer(self, section, fmt, start, first_page_blank=False):
        self._numbering(section, fmt, start)
        section.footer.is_linked_to_previous = False
        if first_page_blank:
            section.different_first_page_header_footer = True
        footer = section.footer
        par = footer.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in list(par.runs):
            run._r.getparent().remove(run._r)
        self._page_field(par)

    # -------------------------------------------------------------- front matter
    def cover(self):
        logo = self.doc.add_paragraph()
        logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo.paragraph_format.space_before = Pt(6)
        logo.paragraph_format.space_after = Pt(10)
        logo.add_run().add_picture(str(ASSETS / "logo.jpg"), width=Inches(1.15))

        self.p("PROJECT REPORT", align=WD_ALIGN_PARAGRAPH.CENTER, size=24, bold=True, space_after=8)
        self.p(
            "A report submitted in partial fulfillment of the requirements for the Award of Degree of",
            align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=11, space_after=8,
        )
        self.p("BACHELOR OF ENGINEERING", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, space_after=4)
        self.p("in", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=4)
        self.p(
            "CSE (ARTIFICIAL INTELLIGENCE AND MACHINE LEARNING)",
            align=WD_ALIGN_PARAGRAPH.CENTER, size=12.5, bold=True, space_after=18,
        )
        self.p(PROJECT_SHORT, align=WD_ALIGN_PARAGRAPH.CENTER, size=30, bold=True, space_after=2)
        self.p(
            "Inventory and Stock Management System",
            align=WD_ALIGN_PARAGRAPH.CENTER, size=15, bold=True, space_after=10,
        )
        self.p(
            "\u201cA single place to know what is in stock, what is moving and what needs replenishing.\u201d",
            align=WD_ALIGN_PARAGRAPH.RIGHT, italic=True, size=10.5, space_after=12,
        )
        self.p("Submitted by", align=WD_ALIGN_PARAGRAPH.CENTER, size=11.5, bold=True, space_after=6)
        for name, roll in TEAM:
            self.p(f"{name}    ({roll})", align=WD_ALIGN_PARAGRAPH.CENTER, size=11.5, bold=True, space_after=2)
        self.p("Under the supervision of", align=WD_ALIGN_PARAGRAPH.CENTER, size=11.5, bold=True,
               space_before=12, space_after=4)
        self.p("Mr. Ravi P", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, bold=True, space_after=2)
        self.p("Assistant Professor, Department of CSE (AIML)", align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, space_after=2)
        self.p("KPR Institute of Engineering and Technology, Coimbatore",
               align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, space_after=18)
        self.p("KPR INSTITUTE OF ENGINEERING AND TECHNOLOGY", align=WD_ALIGN_PARAGRAPH.CENTER,
               size=13, bold=True, space_after=2)
        self.p("AVINASHI ROAD, ARASUR, COIMBATORE \u2013 641 407",
               align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, bold=True, space_after=8)
        self.p("2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, bold=True, space_after=0)

    def certificate(self):
        logo = self.doc.add_paragraph()
        logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo.paragraph_format.space_after = Pt(8)
        logo.add_run().add_picture(str(ASSETS / "logo.jpg"), width=Inches(0.8))
        self.p("KPR INSTITUTE OF ENGINEERING AND TECHNOLOGY", align=WD_ALIGN_PARAGRAPH.CENTER,
               size=14, bold=True, space_after=6)
        title = self.p("BONAFIDE CERTIFICATE", align=WD_ALIGN_PARAGRAPH.CENTER, size=15, bold=True, space_after=18)
        title.runs[0].underline = True
        body = [
            ("This is to certify that the Project report titled ", False),
            ("\u201cStockFlow \u2013 Inventory and Stock Management System\u201d", True),
            (" submitted by SAI VARSHINI J (25AM104), SAPARNA S (25AM105), SASMITHA K M (25AM106) and "
             "SENAAN Y A (25AM107) is work done by them and submitted during the academic year 2026 \u2013 2027, "
             "in partial fulfillment of the requirements for the award of the degree of Bachelor of Engineering "
             "in CSE (Artificial Intelligence and Machine Learning).", False),
        ]
        self.runs(body, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, space_after=24, line_spacing=1.5)

        self.doc.add_paragraph()
        table = self.doc.add_table(rows=2, cols=3)
        table.autofit = False
        widths = [2.1, 2.1, 2.0]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        entries = [
            ("Mentor", "Mr. Ravi P", "Assistant Professor II,\nDepartment of CSE (AIML),\nKPR Institute of Engineering and\nTechnology, Coimbatore."),
            ("Skilling Officer", "Mrs. Sri Sakthi A", "Assistant Professor II,\nDepartment of CSE (AIML),\nKPR Institute of Engineering and\nTechnology, Coimbatore."),
            ("Head of the Department", "Dr. Pandiya Rajan G", "Assistant Professor III,\nDepartment of CSE (AIML),\nKPR Institute of Engineering and\nTechnology, Coimbatore."),
        ]
        for col, (role, name, detail) in enumerate(entries):
            cell = table.rows[0].cells[col]
            cell.width = Inches(widths[col])
            cell.text = role
            for par in cell.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                par.paragraph_format.space_after = Pt(2)
                for run in par.runs:
                    run.bold = True
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM

            cell = table.rows[1].cells[col]
            cell.width = Inches(widths[col])
            cell.text = name + "\n" + detail
            for i, par in enumerate(cell.paragraphs):
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.0
                for run in par.runs:
                    run.bold = i == 0
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.5)

        for row in table.rows:
            for cell in row.cells:
                self._margins(cell, top=30, bottom=30, left=60, right=60)
                tc_pr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    node = OxmlElement(f"w:{edge}")
                    node.set(qn("w:val"), "nil")
                    borders.append(node)
                tc_pr.append(borders)

        self.p("", space_after=24)
        self.runs([("Place: ", True), ("Coimbatore", False)], align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
        self.runs([("Date: ", True), ("October 2026", False)], align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)

    def acknowledgement(self):
        self.p("ACKNOWLEDGEMENT", align=WD_ALIGN_PARAGRAPH.CENTER, size=15, bold=True, space_after=18)
        paragraphs = [
            [("We would like to express our sincere gratitude to ", False),
             ("Dr. K. P. Ramasamy, Chairman", True),
             (", for providing us with the opportunity and necessary facilities to undertake this project.", False)],
            [("We extend our heartfelt thanks to ", False), ("Dr. M. Akila, Principal", True),
             (", for her valuable support, encouragement and guidance throughout the course of this project.", False)],
            [("We would like to express our sincere thanks to ", False),
             ("Dr. Pandiya Rajan G, Head of the Department, CSE (Artificial Intelligence and Machine Learning)", True),
             (", for his valuable support, encouragement and permission to carry out this project.", False)],
            [("We are deeply grateful to ", False), ("Mrs. Sri Sakthi A, Skilling Officer", True),
             (", for her valuable support, encouragement and guidance throughout the project.", False)],
            [("We would like to express our sincere gratitude to our mentor, ", False), ("Mr. Ravi P", True),
             (", Department of CSE (Artificial Intelligence and Machine Learning), for his valuable guidance, "
              "continuous support and constructive suggestions throughout the course of this project.", False)],
            [("We also extend our thanks to all the faculty members of the department, our friends and our "
              "families for their encouragement, cooperation and support. Their willingness to share knowledge "
              "and provide motivation contributed greatly to the successful completion of this project.", False)],
            [("Finally, we express our heartfelt gratitude to everyone who directly or indirectly supported us "
              "in completing this project successfully.", False)],
        ]
        for parts in paragraphs:
            self.runs(parts, space_after=12, line_spacing=1.5)

    def abstract(self):
        self.h1("ABSTRACT")
        self.p(
            "Inventory is the working capital of every trading, retail and manufacturing organisation. "
            "Knowing how much of each item is available, where it is kept, what it cost and when it moved "
            "is essential for billing, purchasing, customer service and financial control. In small and "
            "medium-sized organisations these records are still kept in paper registers or in spreadsheets "
            "that are edited by several people. Such records become inconsistent, do not warn anyone when "
            "an item falls below its minimum level, and cannot answer simple questions such as how many "
            "units of a product were issued last week or which supplier is associated with an item.",
        )
        self.p(
            "This report presents StockFlow, a web-based inventory and stock management system that "
            "centralises the product catalogue, categories, suppliers and every stock movement in one "
            "application. The system provides a dashboard with live totals, low-stock alerts and quick "
            "actions; a searchable product catalogue with minimum and maximum stock levels; forms for "
            "receiving and issuing stock; a complete transaction history; operational reports; a low-stock "
            "queue with suggested reorder quantities; and an administrator-only user management screen. "
            "A React and TypeScript single-page interface communicates with a FastAPI REST API. Business "
            "rules, such as rejecting an issue that exceeds the available quantity, are enforced on the "
            "server, while SQLAlchemy stores users, products and movements in a SQLite database.",
        )
        self.p(
            "The application was verified through eight automated backend tests and a feature-by-feature "
            "walkthrough of the user interface. Product creation and validation, stock-in and stock-out "
            "flows, insufficient-stock rejection, low-stock and reorder calculations, dashboard statistics "
            "and role-based access were all confirmed. The report describes the domain, the limitations "
            "of existing methods, the system design (architecture, data model, stock algorithm and "
            "transaction policy), the implementation of each screen, and the testing results. It closes "
            "with the limitations of the current version and a roadmap for barcode support, multi-warehouse "
            "operation, notifications, forecasting and cloud deployment.",
        )
        self.runs([("Keywords: ", True),
                   ("inventory management, stock control, FastAPI, React, TypeScript, SQLAlchemy, SQLite, "
                    "REST API, role-based access control, audit trail.", False)],
                  space_after=0, line_spacing=1.5)

    # --------------------------------------------------------------- chapters
    def chapter1(self):
        self.h1("Chapter 1  Introduction", page_break=False)
        self.h3("1.1 Background")
        self.p(
            "Every organisation that buys, stores and sells goods depends on accurate stock information. "
            "A retail shop needs to know whether a popular item is available before promising it to a "
            "customer. A warehouse needs to know how many units were received from a supplier and how many "
            "were issued to a department. An accounts team needs the value of the inventory that is held. "
            "When these answers are kept in memory or in a notebook, they are rarely consistent, and the "
            "cost of the resulting mistakes is paid in lost sales, emergency purchases and duplicated orders."
        )
        self.p(
            "The two classic inventory failures pull in opposite directions. Overstocking ties up capital, "
            "fills valuable storage space and increases the risk of damage or obsolescence. Understocking "
            "interrupts sales and production, forces expensive rush orders and damages the reputation of the "
            "business. The balance between the two is managed through minimum and maximum stock levels: the "
            "minimum signals when replenishment must begin, and the maximum caps how much should be held. "
            "A useful system must therefore track quantities continuously and make exceptions visible before "
            "they become shortages."
        )
        self.p(
            "Spreadsheets are the most common digital alternative, and they are flexible, but they were not "
            "designed for transactional data. Several people editing one file create conflicting versions, "
            "formulas can be overwritten accidentally, and there is no reliable record of who changed a "
            "quantity or why. Desktop inventory packages solve part of the problem but often require a "
            "licence, a local installation and a training period. A browser-based application with a clear "
            "data model, server-side validation and an audit trail offers a practical middle ground for "
            "small teams and for academic demonstrations of a complete inventory workflow."
        )
        self.p(
            "StockFlow was built on this idea. It treats every receipt and issue as a transaction, keeps "
            "the current balance on the product record, derives the stock status from the minimum level, "
            "and stores the identity of the user who performed each movement. The interface is deliberately "
            "compact: a dashboard for the daily picture, a catalogue for maintenance, two focused forms for "
            "movement, an audit table and four reports. The result is a system that can be demonstrated and "
            "extended without hiding its logic behind a commercial platform."
        )

        self.h3("1.2 Motivation")
        self.p("The motivation for the project came from the practical difficulties observed in small warehouses and college stores:")
        self.bullets([
            "Poor visibility: staff cannot answer basic questions such as how many units remain, which items are below their minimum level, or how much inventory is held, without manually counting or searching several files.",
            "Weak audit trail: when a quantity changes, there is no reliable record of the previous balance, the new balance, the reference document or the person responsible.",
            "Manual calculations: stock-in and stock-out arithmetic, inventory valuation and reorder quantities are repeated by hand, which introduces avoidable errors.",
            "Fragmented information: product details, supplier contacts and movement records live in separate documents, so cross-checking them is slow.",
            "No role separation: in shared spreadsheets every user can change every value, including prices and user lists, with no restriction.",
            "Reporting effort: producing a low-stock list or a supplier summary requires filtering and reformatting, which is usually postponed until it is needed urgently.",
        ])
        self.p(
            "A second motivation was educational. The project offered an opportunity to connect a modern "
            "frontend, a typed backend and a relational database into one working system, and to apply "
            "software engineering ideas such as layered design, validation, authentication, automated "
            "testing and responsive interface design to a problem that is easy to explain and easy to evaluate."
        )

        self.h3("1.3 Objectives")
        self.p("The objectives of StockFlow are to:")
        self.bullets([
            "provide a single, browser-based workspace for products, categories, suppliers and stock movements;",
            "maintain an accurate current balance for every product and derive its In Stock, Low Stock or Out of Stock status automatically;",
            "record every receipt and issue as an immutable transaction containing the previous balance, the new balance, the quantity, the user, a reference and the time;",
            "reject invalid operations such as a non-positive quantity, a duplicate product code, an issue larger than the available stock, or a maximum level that is not greater than the minimum level;",
            "present a dashboard with inventory totals, low-stock alerts and reorder suggestions so that replenishment can be planned;",
            "generate inventory, stock-movement, low-stock and supplier reports directly from the live database;",
            "separate administrative operations from day-to-day stock operations through role-based access control; and",
            "offer a responsive interface that works on a computer, tablet or mobile browser without installation on the client.",
        ])

        self.h3("1.4 Scope and assumptions")
        self.p("The scope of this version is a single-organisation inventory system that runs on a local network or on one machine for demonstration. The following scope and assumptions apply:")
        self.bullets([
            "The system manages products, categories, suppliers, users, stock-in movements, stock-out movements and reports. It does not manage sales orders, invoices, payments, purchase orders or accounting entries.",
            "Stock is tracked as a single quantity per product. Batch numbers, serial numbers, expiry dates, bin locations and multiple warehouses are outside the current scope.",
            "Authentication is implemented with JSON Web Tokens and bcrypt password hashes. In the delivered demonstration build the interface starts directly as the demo administrator; the login API remains available for API clients.",
            "The backend and the SQLite database run on the same machine. The database file is created automatically on first start and can be populated with the supplied seed script.",
            "It is assumed that users enter correct product, supplier and movement information and that the server clock is correct, because transaction timestamps and report periods depend on it.",
            "The system is a record-keeping and decision-support tool. It does not provide financial advice, tax calculations, medical advice or any other domain-specific recommendation.",
            "Out of scope for this version: barcode scanning, e-mail or SMS notifications, multi-warehouse transfers, demand forecasting, supplier self-service, document attachments and native mobile applications. These are proposed as future work in Chapter 7.",
        ])

    def chapter2(self):
        self.h1("Chapter 2  Domain Review")
        self.h3("2.1 Existing Systems")
        self.p("Organisations currently keep track of stock in one of the following ways:")
        self.p(
            "A notebook or a set of bin cards records receipts and issues. The method is inexpensive and "
            "works without electricity, but it is slow, difficult to total, vulnerable to lost pages and "
            "almost impossible to analyse across products or periods.",
            lead="Manual registers. ",
        )
        self.p(
            "They are the most widespread digital tool: they support formulas, filters and simple charts, "
            "and every office already has one. Their weaknesses appear as soon as more than one "
            "person edits the file: there is no transaction log, no server-side validation, no concurrency "
            "control and no reliable link between a movement and the user who recorded it. A single "
            "overwritten formula can silently corrupt the balance.",
            lead="Spreadsheets. ",
        )
        self.p(
            "Enterprise resource planning and inventory suites provide complete functionality for large "
            "organisations, including purchasing, warehousing, accounting and manufacturing. They are "
            "powerful but expensive, require configuration and training, and impose far more process than a "
            "small store or a college department needs.",
            lead="ERP and commercial inventory suites. ",
        )
        self.p(
            "Many billing applications include a stock module. These tools are convenient for a shop that "
            "already uses them for invoicing, but the inventory features are usually secondary: minimum "
            "levels, reorder suggestions, movement references, role separation and supplier reports may be "
            "missing or locked behind a higher plan.",
            lead="Billing and point-of-sale applications. ",
        )
        self.p(
            "Some teams build a small in-house database or script. This can fit the organisation precisely, "
            "but it creates a maintenance obligation: schema changes, backups, security patches and user "
            "support must be handled internally, and the original author may not be available later.",
            lead="Custom in-house tools. ",
        )
        self.p(
            "The common thread is that each option solves part of the problem. Registers describe but do "
            "not calculate, spreadsheets calculate but do not control, ERP suites control but at a cost and "
            "complexity that small teams cannot justify."
        )

        self.h3("2.2 Stock Movement and Dynamic Inventory")
        self.p(
            "Inventory is not a static number. It changes whenever goods are received, issued, returned, "
            "damaged or counted. A system that stores only a quantity per product cannot explain how that "
            "quantity was reached, so corrections overwrite history and discrepancies cannot be investigated. "
            "StockFlow therefore treats the product record as the current balance and the transaction table "
            "as the ledger. Every operation writes both: the balance is updated and the movement is appended "
            "with the previous and new quantities. The two values can be reconciled at any time."
        )
        self.p(
            "Dynamic behaviour is required in everyday situations. A supplier delivers fewer units than "
            "ordered, a department returns unused items, a physical count reveals damage, or a product price "
            "changes. Table 2.1 summarises how the system responds to the events that occur most often."
        )
        self.table(
            "Table 2.1: Typical inventory events and the response of the system",
            ["Situation", "Response of the system"],
            [
                ["Goods received from a supplier", "The quantity is added to the current balance, the purchase price can be updated, and a STOCK_IN transaction is recorded with the reference and user."],
                ["Items issued to a department or customer", "The requested quantity is checked against the available stock, the balance is reduced, and a STOCK_OUT transaction is recorded with the department in the notes."],
                ["Damaged or expired units", "The units are issued with an explanatory note, so the loss is visible in the audit trail instead of disappearing in an adjustment."],
                ["Physical count correction", "The difference is entered as a stock-in or stock-out movement with a reference such as COUNT-2026-04, preserving the reason."],
                ["Purchase price change", "The product can be edited directly, or the new price can be supplied during a stock-in operation and stored on the product and the transaction."],
                ["Product discontinued", "The product is removed through the admin-only delete action; its dependent transactions are deleted by the configured cascade and the remaining catalogue is unaffected."],
            ],
            [1.85, 4.35],
        )

        self.h3("2.3 Proposed System")
        self.p(
            "StockFlow is a web-based inventory and stock management system with a React single-page "
            "interface and a FastAPI backend. The user works in one browser tab: the sidebar switches "
            "between the dashboard, catalogue, movement forms, audit trail, reports and administration. "
            "The interface calls a versioned REST API, the API validates every request with Pydantic "
            "schemas and business rules, and SQLAlchemy persists the result in SQLite. Because the browser "
            "holds no business data, two users who open the same address see the same records after a "
            "refresh."
        )
        self.p("The main features of the proposed system are:")
        self.bullets([
            "a dashboard with six metric cards (products, units in stock, inventory value, suppliers, low stock and out of stock), a low-stock alert panel and quick actions;",
            "a searchable product catalogue with category and supplier names, minimum and maximum levels, colour-coded status badges, and edit and delete actions for administrators;",
            "maintenance screens for categories and suppliers, including contact, address and GST details;",
            "dedicated Stock In and Stock Out forms that validate quantities, capture references and notes, and update the balance atomically;",
            "a transaction ledger that records the previous balance, the new balance, the user, the type, the quantity and the reference for every movement;",
            "four reports (inventory, stock movement, low stock and suppliers) generated from the live database, plus a low-stock queue with suggested reorder quantities;",
            "administrator-only user management with roles, active status and password changes; and",
            "a responsive layout that collapses the sidebar, reflows cards and allows tables to scroll on small screens.",
        ])

        self.doc.add_page_break()
        self.h3("2.4 Summary of Existing Approaches")
        self.p(
            "Tables 2.2 and 2.3 compare the approaches described above with the proposed system. The "
            "comparison is qualitative: it considers the features that matter for a small inventory "
            "operation rather than the total feature count of a commercial ERP product."
        )
        self.table(
            "Table 2.2: Strengths and limitations of existing approaches",
            ["Approach", "Strength", "Limitation"],
            [
                ["Manual register", "No cost, simple to start", "Slow, no validation, no reports, easily lost"],
                ["Spreadsheet", "Familiar and flexible", "No audit trail, no concurrency, formula errors"],
                ["ERP suite", "Complete and integrated", "Expensive, complex, heavy training and setup"],
                ["Billing application", "Convenient for sales", "Inventory features are secondary or limited"],
                ["Custom in-house tool", "Fits the organisation", "Maintenance and security burden"],
                ["StockFlow (proposed)", "Catalogue, movements, alerts and reports in one place", "Single-organisation, local deployment in this version"],
            ],
            [1.55, 2.35, 2.3],
        )
        self.table(
            "Table 2.3: Feature comparison of existing approaches with the proposed system",
            ["Feature", "Manual register", "Spreadsheet", "ERP suite", "Billing app", "StockFlow"],
            [
                ["Stores product details", "Yes", "Yes", "Yes", "Partial", "Yes"],
                ["Validates every movement", "No", "Partial", "Yes", "Partial", "Yes"],
                ["Minimum / maximum levels", "Manual", "Manual", "Yes", "Rare", "Yes"],
                ["Audit trail with user", "No", "No", "Yes", "Partial", "Yes"],
                ["Role-based access", "No", "No", "Yes", "Partial", "Yes"],
                ["Low-stock and reorder list", "Manual", "Manual", "Yes", "Rare", "Yes"],
                ["Operational reports", "No", "Manual", "Yes", "Yes", "Yes"],
                ["Setup cost and effort", "Very low", "Low", "High", "Medium", "Low"],
            ],
            [1.45, 0.85, 0.9, 0.85, 0.9, 1.25],
            font=9.0,
            center_cols=(1, 2, 3, 4, 5),
        )

    def chapter3(self):
        self.h1("Chapter 3  Gap Identification and Problem Statement")
        self.h3("3.1 The Gap")
        self.p(
            "The review shows that each existing option solves only part of the inventory problem. "
            "Registers and spreadsheets record data but do not enforce rules, warn about shortages or "
            "preserve a trustworthy history. Commercial suites enforce rules but are disproportionate for "
            "a small store, a laboratory or a college department. The specific gaps are:"
        )
        self.bullets([
            "Limited visibility: there is no single view of the current quantity, its value and its status relative to the minimum level.",
            "No reliable audit trail: a changed quantity cannot be traced to a movement, a reference document and a user.",
            "No proactive alerts: low-stock situations are discovered during a count or after a sale has already failed, not when the threshold is crossed.",
            "Error-prone arithmetic: balances and inventory values are calculated manually, and issues larger than the available stock are not always prevented.",
            "Weak access control: shared files allow any user to change prices, delete products or edit the user list.",
            "Fragmented reporting: inventory, movement and supplier information must be assembled from separate files, so routine reporting is delayed.",
            "High cost of adequate tools: the systems that do provide these features are priced and designed for larger organisations.",
        ])

        self.h3("3.2 Problem Statement")
        self.p(
            "Small organisations need to know, at any moment, what stock is available, what it is worth, "
            "what moved and what must be reordered. Manual registers and spreadsheets cannot enforce stock "
            "rules or provide a reliable audit trail, while commercial inventory suites are too expensive "
            "and too complex for the scale of the operation. Hence, a lightweight inventory and stock "
            "management system is required that records products, suppliers and movements in a relational "
            "database, validates every operation on the server, derives stock status and reorder quantities "
            "automatically, presents the information through a clear dashboard and reports, and restricts "
            "administrative actions to authorised users."
        )

        self.h3("3.3 Contribution")
        self.p("The contributions of this work are:")
        self.bullets([
            "a complete, browser-based inventory workspace that combines catalogue maintenance, stock movement, audit history, reports and user administration;",
            "a transaction-centred stock model in which every receipt and issue stores the previous balance, the new balance, the quantity, the reference and the responsible user;",
            "server-side validation of product codes, stock levels, quantities and relationships, so that invalid data cannot enter the database through the interface;",
            "automatic derivation of stock status and reorder quantities from minimum and maximum levels, exposed through dashboard alerts, a low-stock queue and reports;",
            "role-based access control with hashed passwords and signed tokens, plus a self-protection rule that prevents an administrator from deleting, deactivating or demoting their own account; and",
            "a responsive React interface and a documented REST API that can be extended with barcode input, notifications or a different database without rewriting the domain logic.",
        ])
        self.p(
            "Taken together, these contributions address the gaps identified in Section 3.1: the system "
            "makes stock visible, preserves the history of every change, warns before a shortage becomes "
            "critical and limits destructive actions to administrators. The next chapter turns these "
            "ideas into a concrete architecture, data model and set of algorithms."
        )

    def chapter4(self):
        self.h1("Chapter 4  System Design")
        self.h3("4.1 Architecture")
        self.p(
            "StockFlow uses a three-tier architecture. The presentation tier is a React single-page "
            "application that runs in the browser; the application tier is a FastAPI service that owns "
            "the business rules; and the data tier is a SQLite database accessed through SQLAlchemy. "
            "Each tier has one responsibility, and the boundaries are explicit: the browser never sends "
            "SQL, and the server never renders HTML. This separation makes the system easier to test, "
            "because the stock rules can be exercised through the API without a browser, and easier to "
            "replace, because the database can be changed without rewriting the interface."
        )
        self.table(
            "Table 4.1: Layers of the implemented application",
            ["Layer", "Responsibility"],
            [
                ["Presentation layer (React, TypeScript, CSS)", "Dashboard, catalogue, movement forms, audit tables, reports and administration screens; search, modal forms, status badges and responsive layout. Communicates only through the REST API."],
                ["Application layer (FastAPI, Pydantic, services)", "Routing, authentication and role checks, request validation, stock-in and stock-out rules, dashboard aggregation and report generation."],
                ["Data layer (SQLAlchemy, SQLite)", "Persistence of users, categories, suppliers, products and transactions; foreign keys, relationships, cascades and transactions."],
                ["Security layer (JWT, bcrypt, dependencies)", "Password hashing, token creation and verification, active-user checks, admin-only dependencies and the demo-session fallback used by the delivered build."],
            ],
            [1.95, 4.25],
        )
        self.figure("fig_architecture_v2.png", 6.05, "Fig 4.1: Three-tier software architecture of StockFlow")
        self.p(
            "The client tier is served by Vite during development at http://localhost:5173. Axios sends "
            "JSON requests to the API at http://127.0.0.1:8000/api. Every response is either the requested "
            "resource or a JSON object with a human-readable detail message, which the interface displays "
            "in an alert instead of failing silently. A small set of shared functions keeps the screens "
            "consistent: Indian rupee formatting, date formatting, stock-status labels and a request-error "
            "extractor."
        )
        self.p(
            "The application tier is organised by router: authentication, users, categories, suppliers, "
            "products, inventory, dashboard and reports. Routers depend on shared services for stock "
            "movements, dashboard statistics, report generation and database sessions. Pydantic schemas "
            "validate the shape of every request body and query parameter before the handler runs, and "
            "FastAPI converts validation failures into a single readable message. The security dependency "
            "resolves the current user, checks that the account is active, and exposes an admin-only "
            "variant for destructive operations."
        )
        self.p(
            "The data tier is a single SQLite file, stockflow.db, with foreign keys enabled on every "
            "connection. SQLAlchemy models define the columns and relationships, and sessions are created "
            "per request and closed in a finally block. Because SQLite is embedded, the demonstration "
            "needs no database server; because the access layer is SQLAlchemy, the same models can be "
            "moved to PostgreSQL for a multi-user deployment with little change to the service code."
        )

        self.h3("4.2 Data Model")
        self.p(
            "The data model centres on five entities. A category groups products, a supplier provides "
            "them, a user performs movements, and a transaction records each movement of a product. "
            "Table 4.2 lists the attributes that matter for the design; timestamps are present on every "
            "table and are omitted from the descriptions for brevity."
        )
        self.table(
            "Table 4.2: Entities and attributes of the data model",
            ["Entity", "Key attributes", "Description"],
            [
                ["USERS", "id, name, email, password_hash, role, is_active", "Account that signs in, performs movements and, for administrators, manages other users."],
                ["CATEGORIES", "id, name, description", "Logical group used for catalogue navigation and inventory reporting."],
                ["SUPPLIERS", "id, name, contact_person, phone, email, address, city, state, gst_number", "Vendor that supplies products; multiple products may share one supplier."],
                ["PRODUCTS", "id, product_code, name, description, category_id, supplier_id, purchase_price, selling_price, current_stock, minimum_stock, maximum_stock, location", "Master record of an item, including its current balance and replenishment thresholds."],
                ["TRANSACTIONS", "id, product_id, user_id, transaction_type, quantity, previous_stock, new_stock, unit_price, reference_number, notes, created_at", "Immutable ledger entry for every receipt or issue, used by the audit trail and reports."],
            ],
            [1.15, 3.0, 2.05],
            font=9.0,
        )
        self.figure("fig_er_v3.png", 6.05, "Fig 4.2: Entity-relationship model of the StockFlow database")
        self.p(
            "The relationships enforce the operational rules of the domain. One category can contain many "
            "products, and one supplier can provide many products; both foreign keys use ON DELETE RESTRICT, "
            "so a category or supplier that is still referenced cannot be removed by accident. One product "
            "can have many transactions, and the relationship is configured with a cascade so that deleting "
            "a product also removes its movement history; this keeps the ledger consistent with the "
            "catalogue but means that deletion is a destructive, admin-only action. One user can perform "
            "many transactions, and the foreign key uses ON DELETE SET NULL so that historical movements "
            "survive if an account is removed. Stock status is not stored as a column: it is derived from "
            "current_stock and minimum_stock whenever a product is read, which prevents the stored status "
            "and the actual balance from disagreeing."
        )

        self.h3("4.3 Inventory Movement Algorithm")
        self.p(
            "The movement algorithm is the operational core of the system. It receives a product, a "
            "quantity, optional price and reference information, validates the request, updates the "
            "balance and appends a transaction in one database commit. The steps are:"
        )
        self.bullets([
            "Receive the product identifier, the quantity and the optional price, reference and notes.",
            "Load the product; if it does not exist, return a 404 error with a readable message.",
            "For a stock-in movement, reject a quantity that is not greater than zero.",
            "For a stock-out movement, reject a quantity that is not greater than zero and reject a quantity larger than the available balance.",
            "Store the previous balance, calculate the new balance by addition or subtraction, and write it to the product.",
            "For stock-in, update the purchase price when a non-negative unit price is supplied.",
            "Create a transaction with the type, quantity, previous balance, new balance, unit price, reference, notes and the current user.",
            "Commit the product update and the transaction together, so a failure leaves neither change behind.",
            "Refresh the product and return the transaction to the caller.",
            "Let the dashboard, low-stock queue, reports and product status derive their values from the new balance on the next read.",
        ])
        self.p(
            "Flow: select product and operation \u2192 validate quantity \u2192 read previous balance \u2192 "
            "add or subtract \u2192 write transaction \u2192 commit \u2192 refresh dashboard and reports",
            italic=True, size=11, space_after=10,
        )
        self.code([
            "FUNCTION stock_in(product_id, quantity, unit_price):",
            "    product <- FIND Product WHERE id = product_id",
            "    IF product IS NULL THEN RETURN 404",
            "    IF quantity <= 0 THEN RETURN 400",
            "    previous <- product.current_stock",
            "    product.current_stock <- previous + quantity",
            "    IF unit_price IS PROVIDED THEN product.purchase_price <- unit_price",
            "    INSERT Transaction(STOCK_IN, quantity, previous, new_stock, user)",
            "    COMMIT; RETURN transaction",
            "",
            "FUNCTION stock_out(product_id, quantity, reference):",
            "    product <- FIND Product WHERE id = product_id",
            "    IF product IS NULL THEN RETURN 404",
            "    IF quantity <= 0 OR quantity > product.current_stock THEN RETURN 400",
            "    previous <- product.current_stock",
            "    product.current_stock <- previous - quantity",
            "    INSERT Transaction(STOCK_OUT, quantity, previous, new_stock, user)",
            "    COMMIT; RETURN transaction",
            "",
            "FUNCTION status(product):",
            "    IF current_stock = 0 THEN RETURN \"OUT_OF_STOCK\"",
            "    IF current_stock <= minimum_stock THEN RETURN \"LOW_STOCK\"",
            "    RETURN \"IN_STOCK\"",
        ], caption="Algorithm 4.1: Stock movement and status derivation (simplified)")
        self.p(
            "Several rules are enforced before the transaction is written. Product codes are unique and "
            "are stored in upper case, so PRD-KBD-03 and prd-kbd-03 cannot coexist. Prices and quantities "
            "are validated as non-negative, and a product cannot be created or updated with a maximum "
            "level that is not strictly greater than its minimum level. Category and supplier identifiers "
            "must refer to existing rows. The stock-out check compares the requested quantity with the "
            "current balance inside the same request, so an issue that would drive the balance below zero "
            "is rejected with the available quantity in the message. These checks are implemented on the "
            "server, not only in the browser, so they also protect the API when it is called by another client."
        )
        self.p(
            "The initial stock of a new product is handled specially: if a positive initial quantity is "
            "supplied, the product is created and an opening STOCK_IN transaction is written with the "
            "reference INITIAL_OPENING_STOCK. This keeps the audit trail complete from the first unit and "
            "avoids a product whose balance is not explained by any movement."
        )

        self.h3("4.4 Transaction and Event Policy")
        self.p(
            "Inventory changes are event-driven. Every event updates the relevant entity, writes or reuses "
            "an audit entry, and relies on derived values rather than duplicated state. Table 4.3 lists the "
            "principal events and their effect."
        )
        self.table(
            "Table 4.3: Events, state changes and validation",
            ["Event", "Effect on state", "Validation / follow-up"],
            [
                ["Stock received", "Balance increases; STOCK_IN transaction appended", "Quantity > 0; price non-negative; product must exist"],
                ["Stock issued", "Balance decreases; STOCK_OUT transaction appended", "Quantity > 0 and <= available balance"],
                ["Product created with opening stock", "Product inserted; opening STOCK_IN transaction appended", "Unique code; valid category and supplier; max > min"],
                ["Product edited", "Product fields updated; next reads derive a new status", "Unique code if changed; max > min; valid relationships"],
                ["Product deleted", "Product removed; dependent transactions removed by cascade", "Admin only; confirmation in the interface"],
                ["Category or supplier edited", "Related names and report rows reflect the change", "Name uniqueness for categories; valid fields"],
                ["Category or supplier deleted", "Removed only when no product references it", "ON DELETE RESTRICT returns an error otherwise"],
                ["User deactivated", "Account cannot authenticate; historical transactions remain", "Admin only; cannot deactivate own account"],
                ["Page opened or refreshed", "Dashboard, catalogue, alerts and reports are read from the database", "Demo session or bearer token resolves the current user"],
            ],
            [1.55, 2.35, 2.3],
        )
        self.p(
            "The policy has two important consequences. First, a movement and the balance it produces are "
            "committed together, so a crash between the two writes cannot leave the ledger and the product "
            "record out of step. Second, status and reorder quantities are recomputed from the current "
            "balance and the thresholds instead of being stored. This avoids the classic inconsistency in "
            "which a product is marked In Stock while its quantity is zero, and it means that changing the "
            "minimum level immediately changes the alerts and the reports."
        )

        self.h3("4.5 Performance Indicators")
        self.p(
            "The system is judged both as an inventory control tool and as a software product. The "
            "inventory indicators in Table 4.4 can be computed from the product and transaction data, and "
            "they give the organisation a way to evaluate whether the system is improving stock discipline."
        )
        self.table(
            "Table 4.4: Performance indicators",
            ["Indicator", "Formula", "Interpretation"],
            [
                ["Stock accuracy", "Products whose recorded balance equals the physical count / total products x 100", "Higher is better; shows whether the ledger can be trusted"],
                ["Stockout rate", "Products with zero stock / total products x 100", "Lower is better; high values indicate weak replenishment"],
                ["Low-stock ratio", "Products at or below minimum / total products x 100", "Lower is better; reviewed through the low-stock queue"],
                ["Inventory value", "Sum of current_stock x purchase_price over all products", "Working capital held in stock; shown on the dashboard"],
                ["Low-stock response time", "Time between crossing the minimum level and the next stock-in", "Lower is better; measures replenishment discipline"],
                ["Audit completeness", "Transactions with previous balance, new balance and user / total transactions x 100", "Should be 100% in this design"],
            ],
            [1.5, 2.7, 2.0],
            font=9.0,
        )
        self.p(
            "Software performance is measured separately: the latency of the API endpoints, the time "
            "required to build the frontend, the outcome of the automated test suite and the behaviour of "
            "the interface when the backend is unavailable. Chapter 6 reports the measured values and the "
            "limitations of the evaluation."
        )

    def chapter5(self):
        self.h1("Chapter 5  Implementation")
        self.h3("5.1 Technologies Used")
        self.p(
            "The implementation uses a small, well-documented set of technologies. The frontend is a "
            "TypeScript single-page application built with Vite, and the backend is a Python service built "
            "with FastAPI. SQLAlchemy maps the object model to SQLite. This combination keeps the "
            "development loop fast while still exercising the concerns of a real multi-tier application: "
            "HTTP, JSON, validation, authentication, relational constraints and responsive interface design."
        )
        self.table(
            "Table 5.1: Technology stack",
            ["Technology", "Use in the project"],
            [
                ["React 19 + TypeScript", "Component-based interface, typed props and state for products, categories, suppliers, users and transactions."],
                ["Vite 8", "Development server with hot module replacement, production bundling and a proxy for /api requests."],
                ["CSS and Tailwind pipeline", "Custom stylesheet for the application shell, cards, tables, forms, modals and media queries; Tailwind directives remain available for utility styling."],
                ["Axios", "HTTP client for the REST API with centralised error extraction and JSON handling."],
                ["Lucide React", "Consistent outline icons for navigation and actions."],
                ["FastAPI", "REST routers, dependency injection, automatic OpenAPI documentation and exception handling."],
                ["Pydantic 2", "Request and response schemas with field constraints and cross-field validation."],
                ["SQLAlchemy 2.0", "ORM models, sessions, relationships, cascades and transactional writes."],
                ["SQLite", "Embedded relational database stored in backend/stockflow.db; foreign keys are enabled per connection."],
                ["PyJWT + bcrypt", "Signed access tokens and adaptive password hashing."],
                ["Uvicorn", "ASGI server used to run the API on port 8000."],
                ["pytest + HTTPX", "Automated API tests against an in-memory SQLite database."],
            ],
            [1.75, 4.45],
            font=9.0,
        )
        self.p(
            "The interface follows a consistent visual language: a dark navy sidebar, a white content "
            "surface, a blue primary action colour, rounded cards, compact tables and colour-coded status "
            "badges. Green badges mark In Stock, amber badges mark Low Stock and red badges mark Out of "
            "Stock. Every screen uses the same toolbar, spacing and typography, so a user who has learned "
            "the product form can predict how the supplier form behaves. The typography is deliberately "
            "small and dense, because inventory work involves comparing many values at once."
        )

        self.h3("5.2 Application Shell and Dashboard")
        self.p(
            "After the frontend loads, it requests /api/auth/me to resolve the current user. In the "
            "delivered build the API answers with the demo administrator, so the application opens "
            "directly on the dashboard. The shell is a fixed sidebar and a sticky top bar. The sidebar "
            "contains the ten operational screens, the signed-in user and the role; the top bar contains "
            "the menu toggle, the page title, a refresh button and the user avatar. The menu toggle collapses "
            "the sidebar to icons, which is useful on narrow screens and when a wide table needs more room."
        )
        self.p(
            "The dashboard is the landing screen and the daily summary. A gradient hero panel provides a "
            "short greeting and a Record stock action. Six metric cards follow: total products, units in "
            "stock, inventory value at purchase price, suppliers, products at or below their minimum level, "
            "and products with zero stock. The values come from a single dashboard endpoint that aggregates "
            "the product table, so the cards cannot disagree with each other."
        )
        self.p(
            "Below the cards, the low-stock alerts panel lists the products that need attention, ordered by "
            "current quantity, with the product code, minimum level and units remaining. The View all "
            "action opens the full low-stock queue. The quick actions panel provides one-click routes to "
            "Stock In, Stock Out, Products and Transactions. Figure 5.1 shows the dashboard with the seeded "
            "demonstration data: 14 products, 434 units, an inventory value of Rs. 34,119, five suppliers, "
            "four low-stock items and one out-of-stock item."
        )
        self.figure("ui_dashboard.png", 5.9, "Fig 5.1: Dashboard with summary cards, low-stock alerts and quick actions")

        self.h3("5.3 Catalogue Management")
        self.h4("5.3.1 Products List")
        self.p(
            "The Products page is the maintainable view of the catalogue. The table shows the product name "
            "and code, category, supplier, current quantity with the minimum and maximum levels, selling "
            "price, a status badge and, for administrators, edit and delete buttons. The search box filters "
            "by product name or code on the server after a short debounce, so typing does not flood the API. "
            "The page header reports how many products are shown and the list is ordered by creation time, "
            "newest first, which keeps recently added items visible."
        )
        self.p(
            "Status is presented as a badge rather than a raw database value: IN_STOCK becomes a green "
            "In Stock label, LOW_STOCK an amber Low Stock label and OUT_OF_STOCK a red Out of Stock label. "
            "This is the screen where a storekeeper starts the day, because it combines product identity, "
            "availability and price in one row. Figure 5.2 shows the catalogue with its search field and "
            "action column."
        )
        self.figure("ui_products.png", 5.9, "Fig 5.2: Product catalogue with search, stock levels, status badges and row actions")
        self.h4("5.3.2 Add and Edit Product")
        self.p(
            "The product form is presented in a modal over the catalogue, so the user does not lose the "
            "filter or scroll position. It uses a two-column grid for short fields and a full-width "
            "description area. Required fields are enforced by the browser and re-checked by the API. "
            "When editing, the form is pre-filled with the stored values; the initial stock field is shown "
            "only for new products, because later quantity changes must go through the movement forms and "
            "produce an audit entry."
        )
        self.figure("ui_product_form.png", 5.9, "Fig 5.3: Add product modal with catalogue and supplier selection")
        self.table(
            "Table 5.2: Fields of the product form",
            ["Field", "Required", "Purpose and validation"],
            [
                ["Product code", "Yes", "Unique business identifier, 2-50 characters, stored in upper case."],
                ["Name", "Yes", "Descriptive product name, 2-200 characters."],
                ["Category", "Yes", "Must reference an existing category."],
                ["Supplier", "Yes", "Must reference an existing supplier."],
                ["Purchase price", "Yes", "Cost per unit, zero or positive; used for inventory valuation."],
                ["Selling price", "Yes", "Sale price per unit, zero or positive; shown in the catalogue."],
                ["Initial stock", "No (create only)", "Zero or positive; a positive value creates an opening STOCK_IN transaction."],
                ["Minimum stock", "Yes", "Threshold at or below which the product is Low Stock; default 5."],
                ["Maximum stock", "Yes", "Replenishment ceiling and basis for the suggested reorder; must exceed the minimum."],
                ["Location", "No", "Free-text storage location, up to 100 characters."],
                ["Description", "No", "Free-text notes about the item."],
            ],
            [1.45, 1.05, 3.7],
            font=9.0,
        )
        self.p(
            "Saving the form calls POST /api/products or PUT /api/products/{id}. The API rejects a duplicate "
            "code, a missing category or supplier, a negative price, a negative quantity and a maximum level "
            "that is not greater than the minimum. A successful create returns the product with its derived "
            "status and inventory value, and the interface reloads the catalogue and shows a confirmation "
            "toast. Editing works the same way, except that the stock balance is not accepted from the form."
        )
        self.h4("5.3.3 Categories and Suppliers")
        self.p(
            "Categories are managed as a set of cards. Each card shows the category name, its description, "
            "the number of products that use it and, for administrators, an edit action. The product count is "
            "computed with an outer join, so a category with no products still appears with a count of zero. "
            "The API prevents a duplicate name and refuses to delete a category that still contains products."
        )
        self.p(
            "Suppliers are managed in a table with the supplier name, city and state, contact person, phone, "
            "e-mail, product count and an edit action. The form stores the address and GST number in addition "
            "to the contact details. These fields matter for purchasing, because the same supplier may be "
            "contacted about a delivery, a price change or a tax document. Supplier deletion is similarly "
            "blocked while products reference the supplier."
        )

        self.h3("5.4 Stock Operations")
        self.p(
            "Stock In and Stock Out are separate screens because they answer different operational "
            "questions. The Stock In form receives inventory: the user selects a product, enters the "
            "quantity, optionally records the unit purchase price, a reference number and notes, and "
            "confirms. If a purchase price is supplied it is written to the product, so the next inventory "
            "valuation reflects the latest cost. The product selector shows the current quantity for each "
            "item, which helps the user confirm that the correct product was chosen."
        )
        self.figure("ui_stock_in.png", 5.7, "Fig 5.4: Stock In form with product selection, quantity and reference fields")
        self.p(
            "The Stock Out form removes inventory. It validates that the quantity is positive and does not "
            "exceed the current balance, and it offers a customer or department field whose value is "
            "prefixed to the transaction notes. This is a small design decision with a practical benefit: "
            "the audit trail records not only that stock left the store but also where it went, without "
            "requiring a separate issue document. Both forms show a busy state while the request is in "
            "flight, disable the submit button and display the server's error message if the request fails. "
            "On success, the user is taken to the transaction ledger and a toast confirms the operation."
        )
        self.p(
            "The interface does not attempt to calculate the new balance locally. It sends the intent to "
            "the server and displays the result, because the balance may have changed since the page was "
            "loaded. This avoids a stale-browser problem in which two users issue the same stock based on "
            "the same displayed quantity."
        )

        self.h3("5.5 Audit Trail, Reports and Low-Stock Queue")
        self.p(
            "The Transactions page is the audit trail. Each row shows the timestamp, product name and code, "
            "movement type, quantity, the previous and new balance, the user who performed the movement "
            "and the reference. The arrow between the two balances makes the effect of the movement "
            "immediately readable. The search box matches product names, product codes, reference numbers "
            "and notes, so a delivery note or an issue reference can be traced directly. Figure 5.5 shows "
            "the ledger with the seeded movements."
        )
        self.figure("ui_transactions.png", 5.9, "Fig 5.5: Transaction ledger with balances, users and references")
        self.p(
            "The Reports page offers four tabs: Inventory, Stock movement, Low stock and Suppliers. The "
            "inventory report lists products with their category, supplier, quantity and prices; the stock "
            "movement report lists individual movements with their type, quantity and user; the low-stock "
            "report lists products at or below the minimum with a suggested reorder quantity; and the "
            "supplier report aggregates product counts, stock quantities and inventory value by supplier. "
            "Each tab is generated from the live database on selection, and the Refresh action re-runs the "
            "current report. Figure 5.6 shows the inventory report."
        )
        self.figure("ui_reports.png", 5.9, "Fig 5.6: Reports page with inventory, movement, low-stock and supplier tabs")
        self.p(
            "The Low Stock page is a focused replenishment queue. It lists the product, current quantity, "
            "minimum level, suggested reorder quantity, supplier and status, sorted by the most urgent "
            "item. The suggestion is calculated as maximum stock minus current stock, clamped at zero, so "
            "it refills the item towards its ceiling rather than merely to its minimum. Figure 5.7 shows the "
            "queue for the demonstration data."
        )
        self.figure("ui_low_stock.png", 5.7, "Fig 5.7: Low-stock queue with suggested reorder quantities")

        self.h3("5.6 Users and Access Control")
        self.p(
            "The Users page is visible only to administrators. It lists the name, e-mail, role and active "
            "status of each account and allows an administrator to create a user, change a role, reset a "
            "password, deactivate an account or delete it. Passwords are never stored in plain text: bcrypt "
            "produces a salted hash, and the login endpoint compares the submitted password with that hash. "
            "The API issues a signed JSON Web Token containing the user identifier and role; protected "
            "endpoints resolve the token, load the user, reject an inactive account and, for administrative "
            "routes, require the ADMIN role."
        )
        self.figure("ui_users.png", 5.8, "Fig 5.8: Administrator-only user management")
        self.table(
            "Table 5.3: Roles and permissions",
            ["Capability", "ADMIN", "STAFF"],
            [
                ["View dashboard, catalogue, movements and reports", "Yes", "Yes"],
                ["Record stock in and stock out", "Yes", "Yes"],
                ["Create, edit and delete products", "Yes", "No"],
                ["Create and edit categories and suppliers", "Yes", "No"],
                ["Manage users and roles", "Yes", "No"],
            ],
            [4.0, 1.1, 1.1],
            center_cols=(1, 2),
        )
        self.p(
            "Two self-protection rules prevent an administrator from locking the organisation out of its "
            "own administration: an administrator cannot demote or deactivate their own account, and cannot "
            "delete it. The implementation also keeps a demo-session fallback: when a request arrives "
            "without a bearer token, the API resolves the first active administrator, creating the demo "
            "account if the database is empty. This is what allows the delivered interface to open without "
            "a login page. The token flow remains fully functional and is exercised by the automated tests, "
            "but a production deployment should disable the fallback and require sign-in."
        )

        self.h3("5.7 Data Storage and Persistence")
        self.p(
            "All business data is stored in backend/stockflow.db. The file is created automatically on "
            "startup if it does not exist, and SQLAlchemy creates the five tables from the model metadata. "
            "Every connection enables PRAGMA foreign_keys = ON, so the database enforces the relationships "
            "described in Chapter 4 rather than relying on application code alone. Each table carries "
            "created_at and updated_at timestamps; the transaction table carries created_at only, because "
            "a ledger entry is never edited."
        )
        self.p(
            "The supplied seed script populates the database with three users, five categories, five "
            "suppliers, fourteen products and twenty-nine transactions spread over the previous week. The "
            "seed is idempotent: it checks for the demo administrator and exits if the data is already "
            "present, so it can be shipped with a project without duplicating records. The script is also "
            "the quickest way to restore a known demonstration state after experiments."
        )
        self.p(
            "Because SQLite stores the database in a single file, backup is a file copy and the whole "
            "project can be moved between machines. The trade-off is that SQLite allows one writer at a "
            "time; this is acceptable for a small team and for the evaluation workload, but a larger "
            "deployment should move to PostgreSQL. The SQLAlchemy models and most of the service code "
            "would remain unchanged."
        )

        self.h3("5.8 Responsive Interface")
        self.p(
            "The layout is built with CSS grid and flexbox and is adjusted by two media queries. Below "
            "1100 pixels the metric cards change from six columns to three, the dashboard panels stack and "
            "the category cards use two columns. Below 760 pixels the sidebar collapses to a 72-pixel icon "
            "rail, the card grid becomes two columns, the toolbar stacks vertically, the search field takes "
            "the full width and tables scroll horizontally inside their container. The same application "
            "therefore works on a laptop, a tablet and a phone without a separate mobile build. Figure 5.9 "
            "shows the dashboard at a phone width."
        )
        self.figure("ui_mobile.png", 2.05, "Fig 5.9: Dashboard at a phone-width viewport")

        self.h3("5.9 How to Run")
        self.p("The application runs locally with two processes. The backend must be started first so that the dashboard can load its data:")
        self.bullets([
            "Backend: open a terminal in backend, install the requirements with pip install -r requirements.txt, then run python -m uvicorn app.main:app --reload --port 8000. The API is available at http://127.0.0.1:8000 and its interactive documentation at /docs.",
            "Frontend: open a second terminal in frontend, run npm install once, then npm run dev. The application opens at http://localhost:5173.",
            "Windows shortcut: start-backend.bat and start-frontend.bat perform the same steps from the project root.",
            "Demo data: the bundled database is already seeded. To rebuild it, stop the backend, delete backend/stockflow.db, then run python seed.py in the backend folder.",
            "Sign-in: the interface opens as the demo administrator. The login endpoint remains available at POST /api/auth/login for API clients that need a token.",
        ])
        self.p(
            "The backend prints a startup message when the tables are ready, and the frontend prints the "
            "local URL. If the API is stopped, the interface shows a clear message with a Retry button "
            "instead of a blank page."
        )

    def chapter6(self):
        self.h1("Chapter 6  Testing and Evaluation")
        self.h3("6.1 Verification Approach")
        self.p(
            "The system was evaluated at two levels. At the code level, an automated pytest suite "
            "exercises the REST API against an in-memory SQLite database, so each test starts from a known "
            "state and does not depend on the demonstration file. At the interface level, every screen was "
            "walked through feature by feature: records were created, edited, searched and deleted; stock "
            "was received and issued; invalid operations were attempted; and the results were compared with "
            "the expected behaviour. The frontend production build, the TypeScript compiler and the linter "
            "were also run as part of the evaluation."
        )
        self.p(
            "The environment used for the evaluation was an ordinary development machine running Windows "
            "with Python 3.14.5 and Node.js 24.16.0, with the frontend served by Vite 8 and the API by "
            "Uvicorn on loopback ports 5173 and 8000 respectively. The browser was Google Chrome. The "
            "results describe this environment; a different machine or database size may produce different "
            "timings."
        )

        self.h3("6.2 Functional Verification")
        self.p(
            "Table 6.1 lists the functional checks that were performed through the interface and the API. "
            "Each check states what was done, the expected result and the observed outcome."
        )
        self.table(
            "Table 6.1: Functional verification results",
            ["Function", "What was checked", "Expected behaviour", "Result"],
            [
                ["Application startup", "Frontend opened with the backend running and with it stopped", "Dashboard loads; a clear retry message appears when the API is unavailable", "Pass"],
                ["Catalogue", "Products page after seeding", "All products, categories, suppliers and status badges are displayed", "Pass"],
                ["Create product", "New product with category, supplier and opening stock", "Product is saved, opening transaction is created and the row appears", "Pass"],
                ["Duplicate code", "Create or edit with an existing product code", "Request is rejected with a readable message", "Pass"],
                ["Edit product", "Change the name and selling price of a product", "Values are updated and the catalogue reloads in place", "Pass"],
                ["Delete product", "Delete a product that has movements", "Product and its dependent transactions are removed; the page stays on the catalogue", "Pass"],
                ["Stock In", "Receive five units of a test product", "Balance increases from 3 to 8 and a STOCK_IN row is appended", "Pass"],
                ["Stock Out", "Issue five units of the same product", "Balance decreases from 8 to 3 and a STOCK_OUT row is appended", "Pass"],
                ["Insufficient stock", "Issue more units than the available balance", "Request is rejected with the available quantity in the message", "Pass"],
                ["Low-stock queue", "Products at or below the minimum level", "Queue lists the items with current, minimum and suggested reorder quantities", "Pass"],
                ["Dashboard", "Cards after the test movements", "Totals, inventory value and alert counts agree with the database", "Pass"],
                ["Search", "Search products by name and transactions by reference", "Only matching rows are returned", "Pass"],
                ["Reports", "All four report tabs", "Each tab loads live data and returns HTTP 200", "Pass"],
                ["Users", "Create, edit and delete a user as administrator", "Account is updated and the list refreshes without leaving the page", "Pass"],
                ["Persistence", "Reload the browser after changes", "Data is still present and the dashboard reflects the new balances", "Pass"],
            ],
            [1.25, 2.35, 1.95, 0.65],
            font=8.8,
            center_cols=(3,),
        )
        self.p("Result: all functional checks passed in the test environment.")

        self.h3("6.3 Automated Backend Tests")
        self.p(
            "The automated suite contains eight tests. It covers authentication, role enforcement, product "
            "validation, the complete stock-in and stock-out flow, insufficient-stock rejection, low-stock "
            "and reorder calculations, and dashboard statistics. The suite is run with pytest and uses an "
            "in-memory database with a test administrator and a test staff account."
        )
        self.table(
            "Table 6.2: Automated test results",
            ["Test", "Behaviour verified", "Result"],
            [
                ["test_login_success", "Valid credentials return an access token and the user profile", "Pass"],
                ["test_login_invalid_password", "A wrong password returns 401 with a readable message", "Pass"],
                ["test_rbac_admin_only_users_endpoint", "Administrators can list users; staff receive 403", "Pass"],
                ["test_product_crud_and_validation", "Product creation, validation, update and admin-only deletion", "Pass"],
                ["test_stock_in_and_stock_out_flow", "Both movement types update the balance and append transactions", "Pass"],
                ["test_insufficient_stock_rejection", "An issue above the available balance is rejected", "Pass"],
                ["test_low_stock_and_reorder_calculation", "Low-stock alerts and the suggested reorder quantity", "Pass"],
                ["test_dashboard_stats", "Dashboard aggregates are returned with the expected fields", "Pass"],
            ],
            [2.15, 3.35, 0.7],
            font=8.8,
            center_cols=(2,),
        )
        self.p(
            "The final run reported 8 passed, 0 failed and 15 warnings; the warnings are deprecation notices "
            "from the test tooling and the use of timezone-naive timestamps, not failures. The frontend "
            "production build completed successfully (TypeScript and Vite), producing a 305 kB JavaScript "
            "bundle (96 kB gzip) and a 14.8 kB stylesheet (4.2 kB gzip). The linter reported zero errors and "
            "four advisory style warnings."
        )

        self.h3("6.4 Performance Observations")
        self.p(
            "Response times were measured on the loopback interface with a warm process, using fifteen "
            "consecutive requests per endpoint after the data had been seeded. Table 6.3 reports the median "
            "and the 95th percentile. These figures describe the demonstration workload and are not a "
            "benchmark against a production database."
        )
        self.table(
            "Table 6.3: Observed API response times (local machine, 15 requests per endpoint)",
            ["Endpoint", "Median (ms)", "95th percentile (ms)"],
            [
                ["GET /api/health", "0.9", "1.0"],
                ["GET /api/dashboard/stats", "2.2", "2.4"],
                ["GET /api/products (100 rows)", "5.3", "6.1"],
                ["GET /api/inventory/transactions (100 rows)", "6.7", "7.3"],
                ["GET /api/reports/inventory", "4.6", "5.1"],
                ["GET /api/reports/stock-movement", "6.1", "6.6"],
                ["GET /api/reports/low-stock", "2.7", "5.2"],
                ["GET /api/users", "2.2", "2.8"],
            ],
            [3.5, 1.35, 1.35],
            font=9.0,
            center_cols=(1, 2),
        )
        self.p(
            "All endpoints answered within single-digit milliseconds for this data set, and the browser "
            "rendered the resulting tables without a noticeable delay. The most expensive operations are the "
            "transaction and movement reports, because they scan every movement and resolve product and "
            "user names through relationships. The costs are acceptable at the demonstration scale; larger "
            "data sets would benefit from indexes on the date columns and from server-side pagination in the "
            "report endpoints."
        )

        self.h3("6.5 Data Integrity Checks")
        self.p(
            "Several checks specifically targeted the rules that protect the ledger. Deleting a category "
            "that still contains products returned HTTP 400 with a message explaining that the four products "
            "must be reassigned first; deleting a supplier in the same situation returned the equivalent "
            "message. Deleting a test product removed its two movements, confirming the cascade. An invalid "
            "login returned 401, and a staff token was refused access to the users endpoint with 403. "
            "Creating a product with a duplicate code and issuing more stock than was available both "
            "returned validation errors without changing the balance. After each check, the product counts "
            "and transaction totals were read directly from SQLite and matched the interface."
        )

        self.h3("6.6 Threats to Validity and Limitations")
        self.p("The evaluation should be read with the following limitations in mind:")
        self.bullets([
            "Demonstration authentication: the interface opens as the demo administrator, so role restrictions were verified through the API and the test suite rather than by signing in and out of the UI. A production deployment must require authentication and disable the demo fallback.",
            "Single-node database: SQLite serialises writes and the whole database is one file. Concurrent stock operations from many users were not stress-tested.",
            "Limited data volume: the seeded data set is small. Query plans and response times at tens of thousands of transactions were not measured.",
            "No frontend automation: the interface was verified manually. There is no browser test suite to catch regressions in the React components.",
            "Pagination in the interface: the API supports pagination, but the current screens request the first 100 records. A larger catalogue would need page controls.",
            "No notification channel: low-stock information is visible when the application is opened; there are no e-mail, SMS or push alerts.",
            "Single organisation: there is no tenant isolation, and the data model assumes one warehouse and one set of categories.",
            "No concurrency control: two simultaneous issues could both read the same balance before either writes. A production system should use row-level locking or an atomic conditional update.",
            "Functional rather than longitudinal evaluation: the study shows that the features work, not that they improve adherence to inventory policy over months of use.",
        ])

    def chapter7(self):
        self.h1("Chapter 7  Conclusion and Future Work")
        self.p(
            "StockFlow provides a complete, working solution for the inventory problem described in "
            "Chapter 3. It centralises products, categories, suppliers and movements in one application, "
            "keeps a trustworthy balance for every product, records every receipt and issue with the user "
            "and reference, derives stock status and reorder quantities automatically, and presents the "
            "result through a dashboard, an audit ledger and four reports. Administrative actions are "
            "separated from everyday stock work, and the interface adapts to desktop and mobile screens."
        )
        self.p(
            "The objectives stated in Chapter 1 were met. The system records stock-in and stock-out "
            "movements atomically, rejects invalid quantities and duplicate codes, calculates the inventory "
            "value and low-stock counts from live data, produces reorder suggestions from the maximum level, "
            "and keeps the data available between sessions in SQLite. The automated suite and the manual "
            "walkthrough confirm the main flows, and the implementation remains small enough to be "
            "understood and extended by a student team."
        )
        self.p(
            "The evaluation also exposed the boundaries of the current version. The demonstration opens "
            "without a login, the database is embedded and single-writer, the interface fetches at most one "
            "hundred rows per list, and there is no notification channel or concurrency control. These are "
            "deliberate scope choices for a local academic project, not hidden defects, and each has a clear "
            "path in the roadmap below."
        )

        self.h3("Future Work")
        self.p(
            "The current application covers the core inventory workflow. The enhancements in Table 7.1 "
            "would turn it into a more complete operational platform while preserving the same data model "
            "and interface style."
        )
        self.table(
            "Table 7.1: Proposed enhancements",
            ["Enhancement", "Description"],
            [
                ["Barcode and QR support", "Scan a product code during receipt, issue and stock counting to reduce typing and mis-identification."],
                ["Purchase orders", "Raise a purchase order to a supplier, receive it partially or fully, and reconcile the received quantity with the ordered quantity."],
                ["Multi-warehouse operation", "Store balances per location, transfer stock between warehouses and report on each site separately."],
                ["Automated alerts", "Send e-mail or chat notifications when a product crosses its minimum level or when a daily summary is due."],
                ["Cloud deployment", "Move the API and PostgreSQL database to a hosted environment with HTTPS, backups and monitoring."],
                ["Mobile and offline support", "Package the interface as a progressive web application with offline caching for stock counts."],
                ["Demand forecasting", "Use historical movements to estimate future demand and recommend order quantities and reorder points."],
                ["Supplier portal", "Allow suppliers to confirm deliveries, upload invoices and view their own product and performance data."],
                ["Attachments and documents", "Attach invoices, delivery notes and photographs to transactions for a complete documentary trail."],
                ["Granular permissions", "Add roles such as storekeeper, purchase officer and auditor, each with a narrower set of allowed actions."],
                ["Audit export and retention", "Export the ledger to CSV or PDF and define retention and archival rules for historical movements."],
                ["Frontend tests and CI", "Add browser-level tests and a continuous integration pipeline that runs the backend suite, the type checker and the production build on every change."],
            ],
            [1.8, 4.4],
            font=9.0,
            first_col_bold=True,
        )
        self.p(
            "With these enhancements, StockFlow can grow from a well-scoped academic project into a "
            "practical inventory platform for a small organisation, while keeping the clarity, validation "
            "and auditability that make the current version useful."
        )

    def references(self):
        self.h1("References")
        entries = [
            "FastAPI, \u201cFastAPI Documentation\u201d. https://fastapi.tiangolo.com/",
            "React Team, \u201cReact Documentation\u201d. https://react.dev/",
            "Microsoft, \u201cTypeScript Handbook\u201d. https://www.typescriptlang.org/docs/handbook/",
            "Vite, \u201cVite Guide\u201d. https://vite.dev/guide/",
            "SQLAlchemy, \u201cSQLAlchemy 2.0 Documentation\u201d. https://docs.sqlalchemy.org/en/20/",
            "SQLite, \u201cSQLite Documentation\u201d. https://www.sqlite.org/docs.html",
            "Pydantic, \u201cPydantic Documentation\u201d. https://docs.pydantic.dev/",
            "M. Jones, J. Bradley and N. Sakimura, \u201cJSON Web Token (JWT), RFC 7519\u201d, IETF, 2015. https://www.rfc-editor.org/rfc/rfc7519",
            "OWASP, \u201cPassword Storage Cheat Sheet\u201d. https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html",
            "Mozilla, \u201cMDN Web Docs: HTTP\u201d. https://developer.mozilla.org/en-US/docs/Web/HTTP",
            "Mozilla, \u201cMDN Web Docs: Cross-Origin Resource Sharing (CORS)\u201d. https://developer.mozilla.org/en-US/docs/Web/HTTP/CORS",
            "Tailwind Labs, \u201cTailwind CSS Documentation\u201d. https://tailwindcss.com/docs",
            "Lucide, \u201cLucide Icons\u201d. https://lucide.dev/",
            "pytest, \u201cpytest Documentation\u201d. https://docs.pytest.org/",
            "Axios, \u201cAxios Documentation\u201d. https://axios-http.com/docs/intro",
            "W3C, \u201cWeb Content Accessibility Guidelines (WCAG) 2.2\u201d, 2023. https://www.w3.org/TR/WCAG22/",
        ]
        for index, entry in enumerate(entries, 1):
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.35)
            par.paragraph_format.first_line_indent = Inches(-0.35)
            par.paragraph_format.space_after = Pt(4)
            par.paragraph_format.line_spacing = 1.15
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = par.add_run(f"[{index}]  {entry}")
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    def appendix(self):
        self.h1("Appendix A  REST API Summary")
        self.p(
            "The backend exposes the following endpoints under the /api prefix. Token means that the "
            "endpoint accepts a bearer token, or the demo session used by the delivered interface; Admin "
            "means that the ADMIN role is required. The interactive documentation at /docs lists the "
            "request and response schemas and is the authoritative reference; /openapi.json exposes the "
            "same schema for tooling."
        )
        rows = [
            ["Authentication", "POST /auth/login", "Authenticate and return a token", "Public"],
            ["Authentication", "GET /auth/me", "Return the current user", "Token"],
            ["Users", "GET /users", "List accounts", "Admin"],
            ["Users", "POST /users", "Create an account", "Admin"],
            ["Users", "PUT /users/{id}", "Update name, role, status or password", "Admin"],
            ["Users", "DELETE /users/{id}", "Delete an account", "Admin"],
            ["Categories", "GET /categories", "List categories with product counts", "Token"],
            ["Categories", "POST /categories", "Create a category", "Admin"],
            ["Categories", "PUT /categories/{id}", "Update a category", "Admin"],
            ["Categories", "DELETE /categories/{id}", "Delete an unused category", "Admin"],
            ["Suppliers", "GET /suppliers", "List suppliers with product counts", "Token"],
            ["Suppliers", "POST /suppliers", "Create a supplier", "Admin"],
            ["Suppliers", "PUT /suppliers/{id}", "Update a supplier", "Admin"],
            ["Suppliers", "DELETE /suppliers/{id}", "Delete an unused supplier", "Admin"],
            ["Products", "GET /products", "Search, filter, sort and paginate", "Token"],
            ["Products", "GET /products/{id}", "Return one product", "Token"],
            ["Products", "POST /products", "Create a product with optional opening stock", "Admin"],
            ["Products", "PUT /products/{id}", "Update a product", "Admin"],
            ["Products", "DELETE /products/{id}", "Delete a product and its movements", "Admin"],
            ["Inventory", "POST /inventory/stock-in", "Receive stock and append a transaction", "Token"],
            ["Inventory", "POST /inventory/stock-out", "Issue stock and append a transaction", "Token"],
            ["Inventory", "GET /inventory/transactions", "List and search the ledger", "Token"],
            ["Dashboard", "GET /dashboard/stats", "Return the six aggregate values", "Token"],
            ["Dashboard", "GET /dashboard/low-stock-alerts", "Return products below the minimum", "Token"],
            ["Dashboard", "GET /dashboard/reorder-recommendations", "Return suggested order quantities", "Token"],
            ["Reports", "GET /reports/inventory", "Inventory valuation report", "Token"],
            ["Reports", "GET /reports/stock-movement", "Movement report with optional filters", "Token"],
            ["Reports", "GET /reports/low-stock", "Low-stock report with reorder suggestions", "Token"],
            ["Reports", "GET /reports/suppliers", "Supplier summary report", "Token"],
            ["Health", "GET /health", "Service health check", "Public"],
        ]
        self.table(
            "Table A.1: REST endpoints",
            ["Area", "Method and path", "Purpose", "Access"],
            rows,
            [1.15, 2.6, 1.85, 0.6],
            font=8.3,
        )

    def build(self):
        self.cover()
        self.doc.add_page_break()
        self.certificate()
        self.doc.add_page_break()
        self.acknowledgement()
        self.abstract()

        self.doc.add_section(WD_SECTION.NEW_PAGE)
        self.chapter1()
        self.chapter2()
        self.chapter3()
        self.chapter4()
        self.chapter5()
        self.chapter6()
        self.chapter7()
        self.references()
        self.appendix()

        front, body = self.doc.sections[0], self.doc.sections[1]
        self._footer(front, "lowerRoman", 1, first_page_blank=True)
        self._footer(body, "decimal", 1)


def main():
    report = Report()
    report.build()
    report.doc.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    main()
